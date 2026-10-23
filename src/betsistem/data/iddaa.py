"""İddaa oran kayıtları: tek şablon templates/iddaa_kayit_sablonu.xlsx (karar #79).

Kurallar:
- Takım adları football-data yazımıyla girilir (fikstür dosyasındaki HomeTeam/AwayTeam). Her satır,
  `fikstur_dosyasi` sütununda adı yazılan referans fikstür dosyasındaki bir maçla birebir eşleşmek zorundadır.
  Referans dosya önce `betsistem import-fixtures` ile yüklenmiş olmalıdır.
- Kayıt zamanı Türkiye saatidir (+03:00). Kayıt Cuma (hafta içi maçlar için Salı) 17:00–19:00 arasında
  yapılır (#93); aralık dışındaki kayıt saklanır, uyarı verilir ve teste girmez.
- Referans fikstür dosyası, kaydın yapıldığı toplama gününe ait olmalıdır (#95); başka haftanın dosyası hata.
- Dosyada referans oranla (Avg 1-X-2) yer almayan maç, örneğin Cumartesi yüklenen dosyada bulunmayan Cuma
  akşamı maçı, `has_reference = 0` ile saklanır ve teste girmez (#96).
- Dosya birikimli tutulabilir: daha önce eklenmiş satırlar atlanır. Aynı (kayıt zamanı, ev, deplasman) için
  farklı oran veya MBS gelirse hiçbir şey yazılmaz, hata verilir; düzeltme yeni bir kayıt zamanıyla eklenir.
- Hatalı satır varsa dosyanın hiçbir satırı yazılmaz; tüm hatalar birlikte listelenir.
"""

import difflib
import sqlite3
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl

from betsistem.config import LEAGUES, RAW_DIR, season_of
from betsistem.data.fixtures import TR, collection_day, first_upload, fixture_download, fixture_matches
from betsistem.data.raw import RawDataError, store_raw

SOURCE = "iddaa"
SHEET = "Kayıtlar"
COLUMNS = ("kayit_zamani", "lig", "mac_tarihi", "mac_saati", "ev_sahibi", "deplasman",
           "iddaa_1", "iddaa_x", "iddaa_2", "mbs", "fikstur_dosyasi", "not")
REQUIRED = ("kayit_zamani", "lig", "mac_tarihi", "ev_sahibi", "deplasman", "iddaa_1", "iddaa_x", "iddaa_2",
            "mbs", "fikstur_dosyasi")
WINDOW_START, WINDOW_END = time(17, 0), time(19, 0)  # karar #93, Türkiye saati
COLLECTION_WEEKDAYS = (1, 4)  # Salı, Cuma
REFERENCE_COLUMNS = ("AvgH", "AvgD", "AvgA")


@dataclass
class IddaaImportResult:
    added: int = 0
    already_present: int = 0
    warnings: list[str] = field(default_factory=list)


def _parse_datetime(v) -> datetime:
    if isinstance(v, datetime):
        dt = v
    else:
        text = str(v).strip()
        for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S"):
            try:
                dt = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        else:
            raise ValueError(f"kayit_zamani anlaşılamadı: {text!r} (örn. 2026-10-23 19:40)")
    return dt.replace(tzinfo=TR) if dt.tzinfo is None else dt.astimezone(TR)


def _parse_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    text = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"mac_tarihi anlaşılamadı: {text!r} (örn. 2026-10-24)")


def _parse_time(v) -> str | None:
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, (time, datetime)):
        return v.strftime("%H:%M")
    text = str(v).strip()
    try:
        return datetime.strptime(text[:5], "%H:%M").strftime("%H:%M")
    except ValueError as exc:
        raise ValueError(f"mac_saati anlaşılamadı: {text!r} (örn. 17:00)") from exc


def _parse_price(v, name: str) -> float:
    try:
        price = float(str(v).strip().replace(",", ".")) if not isinstance(v, (int, float)) else float(v)
    except ValueError as exc:
        raise ValueError(f"{name} sayı değil: {v!r}") from exc
    if price <= 1.0:
        raise ValueError(f"{name} 1'den büyük olmalı: {price}")
    return price


def _parse_mbs(v) -> int:
    try:
        mbs = int(float(str(v).replace(",", ".")))
    except ValueError as exc:
        raise ValueError(f"mbs tam sayı değil: {v!r}") from exc
    if mbs < 1:
        raise ValueError(f"mbs en az 1 olmalı: {mbs}")
    return mbs


def read_rows(path: Path) -> list[tuple[int, dict]]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if SHEET not in wb.sheetnames:
        raise RawDataError(f"'{SHEET}' sayfası yok; şablon: templates/iddaa_kayit_sablonu.xlsx")
    rows = list(wb[SHEET].iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    if tuple(header[: len(COLUMNS)]) != COLUMNS:
        raise RawDataError(f"Başlık satırı şablonla aynı olmalı: {', '.join(COLUMNS)}")
    out = []
    for i, raw in enumerate(rows[1:], start=2):
        rec = dict(zip(COLUMNS, raw[: len(COLUMNS)]))
        if all(rec[c] in (None, "") for c in COLUMNS):
            continue  # boş satır (yalnızca formül sütunu dolu olabilir)
        out.append((i, rec))
    return out


def _resolve_fixture(conn, name: str):
    info = fixture_download(conn, name)
    if info is None:
        raise ValueError(f"fikstür dosyası '{name}' yüklenmemiş; önce: betsistem import-fixtures {name} --uploaded-at ...")
    first = first_upload(conn, info[1])
    if first is not None and first[0] != name and first[1] < info[2]:
        raise ValueError(
            f"fikstür dosyası '{name}' toplama günü {info[1]} için ilk yükleme değil; referans "
            f"'{first[0]}' (yükleme {first[1]:%d/%m/%y %H:%M} UK) olmalı (#98)"
        )
    return info


def in_collection_window(recorded: datetime) -> bool:
    """#93: Cuma veya Salı, 17:00–19:00 (Türkiye saati, uçlar dahil)."""
    t = recorded.astimezone(TR)
    return t.weekday() in COLLECTION_WEEKDAYS and WINDOW_START <= t.time().replace(second=0, microsecond=0) <= WINDOW_END


def _known_teams(conn, fx, league: str) -> set[str]:
    names = set(fx.loc[fx.Div == league, "HomeTeam"]) | set(fx.loc[fx.Div == league, "AwayTeam"])
    rows = conn.execute(
        "SELECT DISTINCT t.name FROM matches m JOIN teams t ON t.id IN (m.home_team_id, m.away_team_id)"
        " WHERE m.league = ?", (league,)).fetchall()
    return names | {r[0] for r in rows}


def _has_reference_odds(row) -> bool:
    try:
        return all(float(row[c]) > 1.0 for c in REFERENCE_COLUMNS)
    except (KeyError, TypeError, ValueError):
        return False


def import_iddaa(conn: sqlite3.Connection, file: Path, *, raw_dir: Path = RAW_DIR) -> IddaaImportResult:
    file = Path(file)
    rows = read_rows(file)
    result = IddaaImportResult()
    if not rows:
        return result

    errors, parsed = [], []
    fixtures_cache: dict[int, object] = {}
    for line, rec in rows:
        try:
            missing = [c for c in REQUIRED if rec[c] in (None, "")]
            if missing:
                raise ValueError(f"boş zorunlu alan: {', '.join(missing)}")
            recorded = _parse_datetime(rec["kayit_zamani"])
            league = str(rec["lig"]).strip()
            if league not in LEAGUES:
                raise ValueError(f"lig kodu tanınmıyor: {league!r} (örn. E0, T1)")
            match_date = _parse_date(rec["mac_tarihi"])
            if recorded.date() > match_date:
                raise ValueError("kayit_zamani maç tarihinden sonra olamaz")
            home, away = str(rec["ev_sahibi"]).strip(), str(rec["deplasman"]).strip()
            fx_version, fx_day, uploaded = _resolve_fixture(conn, str(rec["fikstur_dosyasi"]).strip())
            if fx_day != collection_day(recorded.date()):
                raise ValueError(
                    f"fikstür dosyası {fx_day} oranlarını taşıyor, kayıt {collection_day(recorded.date())} "
                    "toplama gününe ait; o günün fikstür dosyasını kullan (#95)"
                )
            if fx_version not in fixtures_cache:
                fixtures_cache[fx_version] = fixture_matches(conn, fx_version, raw_dir)
            fx = fixtures_cache[fx_version]
            lg = fx[fx.Div == league]
            hit = lg[(lg.HomeTeam == home) & (lg.AwayTeam == away)]
            has_reference = 1
            if hit.empty:
                names = _known_teams(conn, fx, league)
                unknown = [n for n in (home, away) if n not in names]
                if unknown or match_date > uploaded.date():
                    hints = [f"{n!r} -> {difflib.get_close_matches(n, sorted(names), n=1, cutoff=0.5)}" for n in unknown]
                    raise ValueError(
                        f"{home} - {away} ({league}) referans fikstür dosyasında yok."
                        + (f" Yazım önerisi: {'; '.join(hints)}" if hints
                           else " Maç dosyanın yüklenmesinden sonra oynanıyor; dosyada olmalıydı.")
                    )
                has_reference = 0  # #96: dosya yüklenmeden oynanmış maç (ör. Cuma akşamı)
                result.warnings.append(
                    f"Satır {line}: {home} - {away} ({match_date}) fikstür dosyasında yok (dosya yüklenmeden oynanmış); "
                    "kaydedildi, teste girmez (#96)")
            elif not _has_reference_odds(hit.iloc[0]):
                has_reference = 0
                result.warnings.append(
                    f"Satır {line}: {home} - {away} dosyada var ama referans oranı (Avg 1-X-2) yok; teste girmez (#96)")
            parsed.append({
                "recorded_at": recorded.isoformat(timespec="minutes"),
                "league": league,
                "match_date": match_date.isoformat(),
                "kickoff_time": _parse_time(rec["mac_saati"]),
                "home_team": home,
                "away_team": away,
                "price_home": _parse_price(rec["iddaa_1"], "iddaa_1"),
                "price_draw": _parse_price(rec["iddaa_x"], "iddaa_x"),
                "price_away": _parse_price(rec["iddaa_2"], "iddaa_2"),
                "mbs": _parse_mbs(rec["mbs"]),
                "fixtures_version_id": fx_version,
                "has_reference": has_reference,
                "note": (str(rec["not"]).strip() or None) if rec["not"] is not None else None,
            })
            if not in_collection_window(recorded):
                result.warnings.append(
                    f"Satır {line}: kayıt {recorded:%a %Y-%m-%d %H:%M} Cuma/Salı 17:00–19:00 aralığının dışında; "
                    "kaydedildi, teste girmez (#93)"
                )
        except ValueError as exc:
            errors.append(f"Satır {line}: {exc}")
    if errors:
        raise RawDataError("İddaa dosyasında hata var, hiçbir satır yazılmadı:\n  " + "\n  ".join(errors))

    content = file.read_bytes()
    version_id = store_raw(
        conn, content, source=SOURCE, league="ALL", season=season_of(date.fromisoformat(parsed[0]["match_date"])),
        origin=f"manual:{file.name}", suffix=".xlsx", raw_dir=raw_dir,
    )
    cols = list(parsed[0]) + ["data_version_id"]
    with conn:
        for r in parsed:
            existing = conn.execute(
                "SELECT price_home, price_draw, price_away, mbs FROM iddaa_odds_log"
                " WHERE recorded_at = ? AND home_team = ? AND away_team = ?",
                (r["recorded_at"], r["home_team"], r["away_team"]),
            ).fetchone()
            if existing:
                if tuple(existing) != (r["price_home"], r["price_draw"], r["price_away"], r["mbs"]):
                    raise RawDataError(
                        f"{r['recorded_at']} {r['home_team']} - {r['away_team']} daha önce farklı oran/MBS ile"
                        " kaydedilmiş; kayıtlar değiştirilemez. Düzeltme gerekiyorsa yeni kayit_zamani ile ekle."
                    )
                result.already_present += 1
                continue
            conn.execute(
                f"INSERT INTO iddaa_odds_log ({', '.join(cols)}) VALUES ({', '.join('?' for _ in cols)})",
                (*r.values(), version_id),
            )
            result.added += 1
    return result
